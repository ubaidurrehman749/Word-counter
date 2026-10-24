import base64
import html as htmllib
import io
import json
import re
from collections import Counter
from pathlib import Path

import pandas as pd
import streamlit as st

st.set_page_config(page_title="Word Counter · Study Sparks AI", page_icon="✦", layout="wide")

OWNER_NAME = "Ubaid ur Rehman"


# ---------- Logo ----------
def logo_tag():
    mimes = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
             ".webp": "image/webp", ".svg": "image/svg+xml"}
    for ext, mime in mimes.items():
        p = Path(__file__).parent / f"logo{ext}"
        if p.exists():
            b64 = base64.b64encode(p.read_bytes()).decode()
            return f'<img class="logo" src="data:{mime};base64,{b64}" alt="Study Sparks AI logo">'
    return '<span class="spark">✦</span>'


# ---------- Styling ----------
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Bricolage+Grotesque:wght@600;800&family=DM+Sans:wght@400;500;700&display=swap');
html, body, [class*="css"] { font-family: 'DM Sans', sans-serif; }
#MainMenu, footer, header { visibility: hidden; }
.block-container { padding-top: 1.2rem; max-width: 1100px; }
.topbar { display:flex; justify-content:space-between; align-items:center; flex-wrap:wrap; gap:12px;
          padding-bottom:14px; border-bottom:1px solid #F6D9E1; }
.brand { display:flex; align-items:center; gap:12px; font-family:'Bricolage Grotesque',sans-serif;
         font-weight:800; font-size:1.45rem; color:#1F2433; }
.brand .logo { height:46px; width:46px; object-fit:contain; border-radius:10px; }
.brand .spark { color:#FFC93C; font-size:1.6rem; }
.brand b { color:#D6336C; }
.owner { background:#1F2433; color:#FFF5F7; padding:8px 16px; border-radius:999px; font-weight:500; font-size:.95rem; }
.owner b { color:#FFC93C; }
.hero h1 { font-family:'Bricolage Grotesque',sans-serif; font-weight:800; font-size:2.4rem;
           line-height:1.1; color:#1F2433; margin:1.2rem 0 .3rem; letter-spacing:-.02em; }
.hero p { color:#5B6275; font-size:1.1rem; margin-bottom:.8rem; }

.bigstats { display:grid; grid-template-columns:repeat(4,1fr); gap:12px; margin:6px 0 16px; }
.big { background:#fff; border:1px solid #F6D9E1; border-top:5px solid #D6336C; border-radius:16px;
       padding:16px 12px; text-align:center; box-shadow:0 16px 28px -24px rgba(31,36,51,.55); }
.big:nth-child(2) { border-top-color:#FFC93C; } .big:nth-child(3) { border-top-color:#1F2433; }
.big:nth-child(4) { border-top-color:#2FB883; }
.big b { display:block; font-family:'Bricolage Grotesque',sans-serif; font-size:2rem; color:#1F2433; line-height:1.1; }
.big span { color:#5B6275; font-size:.85rem; font-weight:500; }
@media (max-width:700px){ .bigstats { grid-template-columns:repeat(2,1fr); } }

.row { display:grid; grid-template-columns:repeat(3,1fr); gap:10px; margin-bottom:10px; }
.cell { background:#fff; border:1px solid #F6D9E1; border-radius:12px; padding:10px 14px; }
.cell span { color:#5B6275; font-size:.82rem; display:block; }
.cell b { color:#1F2433; font-size:1.1rem; }
@media (max-width:700px){ .row { grid-template-columns:repeat(2,1fr); } }

div.stButton > button, div.stDownloadButton > button { border-radius:12px; font-weight:700; padding:.6rem 1rem; }
div.stButton > button[kind="primary"], div.stDownloadButton > button {
    background:#D6336C; color:#fff; border:0; }
div.stButton > button[kind="primary"]:hover, div.stDownloadButton > button:hover { background:#B8285A; color:#fff; }
textarea { border-radius:14px !important; }
.stTabs [data-baseweb="tab"] { font-weight:700; }
</style>
""", unsafe_allow_html=True)

st.markdown(f"""
<div class="topbar">
  <div class="brand">{logo_tag()}<span>Study Sparks <b>AI</b></span></div>
  <div class="owner">Created by <b>{OWNER_NAME}</b></div>
</div>
<div class="hero">
  <h1>Word Counter</h1>
  <p>Count words, characters, sentences and reading time. Check keywords and your essay word limit.</p>
</div>
""", unsafe_allow_html=True)

# ---------- File reading ----------
FILE_TYPES = ["pdf", "docx", "pptx", "xlsx", "xlsm", "csv", "txt", "md", "html", "htm",
              "rtf", "json", "png", "jpg", "jpeg", "webp", "bmp", "tiff"]


@st.cache_data(show_spinner=False)
def extract_text(name: str, data: bytes) -> str:
    ext = Path(name).suffix.lower()

    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)

    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" ".join(c.text for c in row.cells))
        return "\n".join(parts)

    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                    parts.append(shape.text_frame.text)
            if slide.has_notes_slide:
                parts.append(slide.notes_slide.notes_text_frame.text)
        return "\n".join(parts)

    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"{ws.title}.")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    parts.append(", ".join(cells) + ".")
        return "\n".join(parts)

    if ext in (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff"):
        from PIL import Image
        import pytesseract
        return pytesseract.image_to_string(Image.open(io.BytesIO(data)))

    text = data.decode("utf-8", errors="ignore")

    if ext in (".html", ".htm"):
        text = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", text)
        text = re.sub(r"(?s)<[^>]+>", " ", text)
        return htmllib.unescape(text)

    if ext == ".rtf":
        text = re.sub(r"\\'[0-9a-f]{2}", " ", text)
        text = re.sub(r"\\[a-z]+-?\d* ?", " ", text)
        return re.sub(r"[{}]", " ", text)

    if ext == ".json":
        try:
            vals = []

            def walk(o):
                if isinstance(o, dict):
                    for v in o.values():
                        walk(v)
                elif isinstance(o, list):
                    for v in o:
                        walk(v)
                elif isinstance(o, str):
                    vals.append(o)
            walk(json.loads(text))
            return "\n".join(vals)
        except Exception:
            return text

    return text


# ---------- Counting logic ----------
STOP = set("""a about above after again all also am an and any are as at be because been before being below
between both but by can could did do does doing down during each few for from further had has have having he her
here hers him his how i if in into is it its just me more most my no nor not now of off on once only or other our
out over own same she should so some such than that the their them then there these they this those through to too
under until up very was we were what when where which while who whom why will with would you your""".split())


def words_of(text):
    return re.findall(r"[\w'’-]+", text)


def syllables(word):
    w = re.sub(r"[^a-z]", "", word.lower())
    if not w:
        return 1
    count = len(re.findall(r"[aeiouy]+", w))
    if w.endswith("e") and not w.endswith("le") and count > 1:
        count -= 1
    return max(1, count)


def fmt_time(minutes):
    secs = round(minutes * 60)
    if secs < 60:
        return f"{secs} sec"
    m, s = divmod(secs, 60)
    return f"{m} min {s} sec" if s else f"{m} min"


def analyze(text):
    words = words_of(text)
    n = len(words)
    sentences = [s for s in re.split(r"(?<=[.!?۔؟])\s+", text.strip()) if words_of(s)]
    paragraphs = [p for p in re.split(r"\n\s*\n", text) if p.strip()]
    lines = [l for l in text.splitlines() if l.strip()]
    ns = max(len(sentences), 1)
    r = {
        "words": n,
        "chars": len(text),
        "chars_ns": len(re.sub(r"\s", "", text)),
        "sentences": len(sentences),
        "paragraphs": len(paragraphs),
        "lines": len(lines),
        "unique": len({w.lower() for w in words}),
        "avg_word": (sum(len(w) for w in words) / n) if n else 0,
        "avg_sent": n / ns if n else 0,
        "longest": max(words, key=len) if words else "-",
        "read": n / 220,
        "speak": n / 130,
        "flesch": None,
    }
    if n >= 30 and sentences:
        syl = sum(syllables(w) for w in words)
        r["flesch"] = 206.835 - 1.015 * (n / ns) - 84.6 * (syl / n)
    return r


def flesch_label(score):
    if score >= 90: return "Very easy"
    if score >= 70: return "Easy"
    if score >= 60: return "Standard"
    if score >= 50: return "Fairly difficult"
    if score >= 30: return "Difficult"
    return "Very difficult"


def keyword_table(text, top, min_len, skip_stop):
    words = [w.lower() for w in words_of(text)]
    total = max(len(words), 1)
    f = Counter(w for w in words
                if len(w) >= min_len and not w.isdigit() and not (skip_stop and w in STOP))
    return pd.DataFrame([{"Word": w, "Count": c, "Density %": round(c / total * 100, 2)}
                         for w, c in f.most_common(top)])


SAMPLE = ("Photosynthesis is the process by which green plants, algae and some bacteria convert light energy into "
          "chemical energy. It takes place mainly in the chloroplasts, which contain a green pigment called "
          "chlorophyll. Chlorophyll absorbs sunlight, especially in the red and blue wavelengths, and reflects "
          "green light, which is why leaves look green.\n\nThe process has two main stages. In the light-dependent "
          "reactions, sunlight splits water molecules, releasing oxygen and producing the energy carriers ATP and "
          "NADPH. In the Calvin cycle, the plant uses ATP and NADPH to turn carbon dioxide from the air into "
          "glucose. Photosynthesis is important because it produces the oxygen we breathe and forms the base of "
          "nearly every food chain.")

# ---------- State ----------
ss = st.session_state
ss.setdefault("text", "")
ss.setdefault("up_key", 0)


def load_sample():
    ss.text = SAMPLE


def clear_all():
    ss.text = ""
    ss.up_key += 1


def use_as_text(value):
    ss.text = value


tab_count, tab_kw, tab_tools = st.tabs(["① Count", "② Keywords", "③ Text tools"])

# ======================= COUNT =======================
with tab_count:
    st.text_area("Paste text", key="text", height=220,
                 placeholder="Type or paste your text here. Press Ctrl+Enter (or tap outside the box) to update the count.",
                 label_visibility="collapsed")
    files = st.file_uploader("Or upload files (PDF, Word, PowerPoint, Excel, CSV, TXT, images and more)",
                             type=FILE_TYPES, accept_multiple_files=True, key=f"up_{ss.up_key}")

    extracted = []
    for f in files or []:
        try:
            with st.spinner(f"Reading {f.name}…"):
                t = extract_text(f.name, f.getvalue())
            if t.strip():
                extracted.append(t)
                st.caption(f"✅ {f.name}: {len(words_of(t)):,} words found")
            else:
                st.warning(f"{f.name}: no readable text found. For scanned PDFs, upload pages as images.")
        except Exception as e:
            st.warning(f"{f.name}: could not be read ({e})")

    full_text = "\n\n".join([ss.text] + extracted).strip()

    b1, b2, _ = st.columns([1, 1, 3])
    b1.button("Try sample", on_click=load_sample, use_container_width=True)
    b2.button("Clear", on_click=clear_all, use_container_width=True)

    R = analyze(full_text)

    st.markdown(f"""
    <div class="bigstats">
      <div class="big"><b>{R['words']:,}</b><span>Words</span></div>
      <div class="big"><b>{R['chars']:,}</b><span>Characters</span></div>
      <div class="big"><b>{R['sentences']:,}</b><span>Sentences</span></div>
      <div class="big"><b>{R['paragraphs']:,}</b><span>Paragraphs</span></div>
    </div>
    <div class="row">
      <div class="cell"><span>Characters (no spaces)</span><b>{R['chars_ns']:,}</b></div>
      <div class="cell"><span>Lines</span><b>{R['lines']:,}</b></div>
      <div class="cell"><span>Unique words</span><b>{R['unique']:,}</b></div>
      <div class="cell"><span>Reading time</span><b>{fmt_time(R['read'])}</b></div>
      <div class="cell"><span>Speaking time</span><b>{fmt_time(R['speak'])}</b></div>
      <div class="cell"><span>Average word length</span><b>{R['avg_word']:.1f} letters</b></div>
      <div class="cell"><span>Average sentence length</span><b>{R['avg_sent']:.1f} words</b></div>
      <div class="cell"><span>Longest word</span><b>{htmllib.escape(R['longest'])}</b></div>
      <div class="cell"><span>Readability (English)</span><b>{
          (f"{R['flesch']:.0f}/100 · " + flesch_label(R['flesch'])) if R['flesch'] is not None else "Needs 30+ words"
      }</b></div>
    </div>
    """, unsafe_allow_html=True)

    st.markdown("#### Word limit")
    g1, g2 = st.columns([1, 2])
    target = g1.number_input("Target words", min_value=0, value=500, step=50)
    mode = g2.radio("Rule", ["Stay under the limit", "Reach at least"], horizontal=True)
    if target:
        st.progress(min(R["words"] / target, 1.0))
        diff = target - R["words"]
        if mode.startswith("Stay"):
            if diff >= 0:
                st.success(f"{R['words']:,} of {target:,} words. {diff:,} words left.")
            else:
                st.error(f"{-diff:,} words over the limit. Trim your text.")
        else:
            if diff > 0:
                st.warning(f"{R['words']:,} of {target:,} words. {diff:,} more words needed.")
            else:
                st.success(f"Target reached: {R['words']:,} words.")

    if R["words"]:
        report = (f"Study Sparks AI · Word Counter report\n{'-' * 40}\n"
                  f"Words: {R['words']}\nCharacters: {R['chars']}\nCharacters (no spaces): {R['chars_ns']}\n"
                  f"Sentences: {R['sentences']}\nParagraphs: {R['paragraphs']}\nLines: {R['lines']}\n"
                  f"Unique words: {R['unique']}\nReading time: {fmt_time(R['read'])}\n"
                  f"Speaking time: {fmt_time(R['speak'])}\nAverage word length: {R['avg_word']:.1f}\n"
                  f"Average sentence length: {R['avg_sent']:.1f}\nLongest word: {R['longest']}\n")
        st.download_button("Download report (.txt)", report, file_name="word-count-report.txt")

# ======================= KEYWORDS =======================
with tab_kw:
    if R["words"] < 5:
        st.info("Add some text in the **Count** tab to see the most used words.")
    else:
        k1, k2, k3 = st.columns(3)
        top = k1.slider("Words to show", 5, 30, 12)
        min_len = k2.slider("Minimum word length", 1, 8, 3)
        skip = k3.checkbox("Ignore common words (the, and, is…)", value=True)
        df = keyword_table(full_text, top, min_len, skip)
        if df.empty:
            st.warning("No words match these settings. Lower the minimum length.")
        else:
            st.bar_chart(df.set_index("Word")["Count"], color="#D6336C")
            st.dataframe(df, use_container_width=True, hide_index=True)
            st.download_button("Download keywords (.csv)", df.to_csv(index=False),
                               file_name="keywords.csv", mime="text/csv")

# ======================= TOOLS =======================
def sentence_case(t):
    t = t.lower()
    return re.sub(r"(^\s*|[.!?]\s+)([a-z])", lambda m: m.group(1) + m.group(2).upper(), t)


TOOLS = {
    "UPPERCASE": lambda t: t.upper(),
    "lowercase": lambda t: t.lower(),
    "Title Case": lambda t: t.title(),
    "Sentence case": sentence_case,
    "Remove extra spaces": lambda t: "\n".join(re.sub(r"[ \t]+", " ", l).strip() for l in t.splitlines()),
    "Remove empty lines": lambda t: "\n".join(l for l in t.splitlines() if l.strip()),
    "Remove duplicate lines": lambda t: "\n".join(dict.fromkeys(t.splitlines())),
    "Sort lines A to Z": lambda t: "\n".join(sorted(t.splitlines(), key=str.lower)),
}

with tab_tools:
    if not full_text:
        st.info("Add some text in the **Count** tab first.")
    else:
        op = st.selectbox("Choose a tool", list(TOOLS))
        result = TOOLS[op](full_text)
        st.text_area("Result", value=result, height=220, disabled=True)
        t1, t2 = st.columns(2)
        t1.button("Use this as my text", on_click=use_as_text, args=(result,), type="primary",
                  use_container_width=True)
        t2.download_button("Download result (.txt)", result, file_name="result.txt", use_container_width=True)
        st.caption("Tip: select the result text and copy it, or download it as a file.")

st.markdown(f"<hr><center style='color:#5B6275'>© Study Sparks AI · Created by {OWNER_NAME}</center>",
            unsafe_allow_html=True)
